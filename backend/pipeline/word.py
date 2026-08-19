#!/usr/bin/env python3
"""把測試報告 JSON 填進參考 Word 範本的 ADO 區塊。"""

from __future__ import annotations

import json
import re
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

from backend.utils.paths import ADO_DIR, DOCX_DIR, TEMPLATE_PATH
from backend.utils.pipeline_io import log, log_stage, log_stage_done, normalize_ids, read_pipeline_json, run_pipeline_main, write_pipeline_json
from backend.utils.report_schema import META_FIELDS, find_report


IMAGE_WIDTH = Inches(6.2)
FONT_NAME = "Microsoft YaHei"
FONT_SIZE = "16"
META_FONT_SIZE = "20"
META_NO_START = 1


def load_report(work_item_id: int) -> dict[str, Any]:
    report = find_report(work_item_id)
    if report is None:
        raise FileNotFoundError(
            f"找不到包含 ADO #{work_item_id} 的 JSON。請先執行 python3 report_by_llm.py"
        )
    return report


def elm_text(elm: etree._Element) -> str:
    return "".join(node.text or "" for node in elm.iter(qn("w:t")))


def set_paragraph_text(p_elm: etree._Element, text: str) -> None:
    nodes = list(p_elm.iter(qn("w:t")))
    if not nodes:
        p = p_elm
        r = p.makeelement(qn("w:r"), {})
        t = r.makeelement(qn("w:t"), {})
        t.text = text
        p.append(r)
        return
    nodes[0].text = text
    if text.startswith(" ") or text.endswith(" "):
        nodes[0].set(qn("xml:space"), "preserve")
    for node in nodes[1:]:
        node.text = ""


def ensure_child(parent: etree._Element, tag: str) -> etree._Element:
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def apply_run_font(r_pr: etree._Element, font_size: str | None = None) -> None:
    r_fonts = ensure_child(r_pr, "w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_NAME)
    r_fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_fonts.set(qn("w:eastAsia"), FONT_NAME)
    r_fonts.set(qn("w:cs"), FONT_NAME)
    r_fonts.set(qn("w:hint"), "eastAsia")
    if font_size is not None:
        ensure_child(r_pr, "w:sz").set(qn("w:val"), font_size)
        ensure_child(r_pr, "w:szCs").set(qn("w:val"), font_size)
    lang = ensure_child(r_pr, "w:lang")
    lang.set(qn("w:eastAsia"), "zh-TW")


def clean_cell_format(cell: Any, align: str = "left", font_size: str = FONT_SIZE) -> None:
    tc_pr = ensure_child(cell._tc, "w:tcPr")
    no_wrap = tc_pr.find(qn("w:noWrap"))
    if no_wrap is not None:
        tc_pr.remove(no_wrap)
    ensure_child(tc_pr, "w:vAlign").set(qn("w:val"), "center")

    for paragraph in cell.paragraphs:
        p_pr = ensure_child(paragraph._p, "w:pPr")
        p_style = p_pr.find(qn("w:pStyle"))
        if p_style is not None:
            p_pr.remove(p_style)
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            p_pr.remove(num_pr)
        indent = p_pr.find(qn("w:ind"))
        if indent is not None:
            p_pr.remove(indent)
        ensure_child(p_pr, "w:jc").set(qn("w:val"), align)
        apply_run_font(ensure_child(p_pr, "w:rPr"), font_size=font_size)


def set_cell_text(
    cell: Any,
    text: str,
    align: str = "left",
    font_size: str = FONT_SIZE,
) -> None:
    lines = [line for line in str(text or "").split("\n")]
    if not lines:
        lines = [""]
    while len(cell.paragraphs) > 1:
        cell.paragraphs[-1]._element.getparent().remove(cell.paragraphs[-1]._element)

    def write_paragraph(paragraph: Any, line: str) -> None:
        proto = deepcopy(paragraph.runs[0]._r) if paragraph.runs else None
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)
        if proto is None:
            run_elm = OxmlElement("w:r")
            paragraph._p.append(run_elm)
        else:
            run_elm = deepcopy(proto)
            for child in list(run_elm):
                if child.tag in {qn("w:t"), qn("w:br"), qn("w:tab")}:
                    run_elm.remove(child)
            paragraph._p.append(run_elm)
        r_pr = ensure_child(run_elm, "w:rPr")
        apply_run_font(r_pr, font_size=font_size)
        t_elm = OxmlElement("w:t")
        t_elm.text = line
        if line[:1].isspace() or line[-1:].isspace():
            t_elm.set(qn("xml:space"), "preserve")
        run_elm.append(t_elm)

    write_paragraph(cell.paragraphs[0], lines[0])
    for line in lines[1:]:
        new_p = deepcopy(cell.paragraphs[0]._p)
        for run in list(new_p.findall(qn("w:r"))):
            new_p.remove(run)
        cell._tc.append(new_p)
        write_paragraph(cell.paragraphs[-1], line)
    clean_cell_format(cell, align=align, font_size=font_size)


def find_paragraph(body: etree._Element, predicate) -> etree._Element | None:
    for child in list(body):
        if child.tag == qn("w:p") and predicate(elm_text(child)):
            return child
    return None


def ensure_row_count(table: Any, total_rows: int) -> None:
    tbl = table._tbl
    while len(table.rows) < total_rows:
        tbl.append(deepcopy(tbl.tr_lst[-1]))
    while len(table.rows) > total_rows:
        tbl.remove(tbl.tr_lst[-1])


def fill_meta_table(table: Any, report: dict[str, Any]) -> None:
    ensure_row_count(table, 1 + len(META_FIELDS))
    header_row = table.rows[0]
    set_cell_text(header_row.cells[0], "No.", align="center", font_size=META_FONT_SIZE)
    set_cell_text(header_row.cells[1], "項目", align="center", font_size=META_FONT_SIZE)
    set_cell_text(header_row.cells[2], "描述", align="center", font_size=META_FONT_SIZE)
    for index, (label, key) in enumerate(META_FIELDS, start=1):
        row = table.rows[index]
        set_cell_text(
            row.cells[0],
            str(META_NO_START + index - 1),
            align="center",
            font_size=META_FONT_SIZE,
        )
        set_cell_text(row.cells[1], label, align="left", font_size=META_FONT_SIZE)
        set_cell_text(
            row.cells[2],
            str(report.get(key) or ""),
            align="left",
            font_size=META_FONT_SIZE,
        )


def set_tbl_width(tbl: etree._Element, widths: list[int]) -> None:
    """固定表格與欄寬（dxa），並取消負縮排。"""
    total = sum(widths)
    tbl_pr = ensure_child(tbl, "w:tblPr")
    tbl_w = ensure_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is not None:
        tbl_ind.set(qn("w:w"), "0")
        tbl_ind.set(qn("w:type"), "dxa")
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(1, grid)
    cols = list(grid.findall(qn("w:gridCol")))
    while len(cols) < len(widths):
        col = OxmlElement("w:gridCol")
        grid.append(col)
        cols.append(col)
    for col in cols[len(widths) :]:
        grid.remove(col)
    for col, width in zip(cols, widths):
        col.set(qn("w:w"), str(width))

    for row in tbl.findall(qn("w:tr")):
        cells = row.findall(qn("w:tc"))
        for cell, width in zip(cells, widths):
            tc_pr = ensure_child(cell, "w:tcPr")
            tc_w = ensure_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def grid_widths(tbl: etree._Element) -> list[int]:
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        return []
    widths: list[int] = []
    for col in grid.findall(qn("w:gridCol")):
        raw = col.get(qn("w:w")) or "0"
        widths.append(int(raw) if str(raw).isdigit() else 0)
    return widths


def scaled_widths(source: list[int], target_total: int) -> list[int]:
    current = sum(source) or 1
    scaled = [max(1, round(width * target_total / current)) for width in source]
    scaled[-1] += target_total - sum(scaled)
    return scaled


def fill_items_table(table: Any, items: list[dict[str, Any]], *, target_width: int | None = None) -> None:
    rows_needed = max(2, 1 + len(items))
    ensure_row_count(table, rows_needed)
    if not items:
        set_cell_text(table.rows[1].cells[0], "1", align="center")
        for col in range(1, 5):
            set_cell_text(table.rows[1].cells[col], "", align="left")
    else:
        while len(table.rows) > 1 + len(items):
            table._tbl.remove(table.rows[-1]._tr)
        for index, item in enumerate(items, start=1):
            row = table.rows[index]
            data = item.get("test_data") or []
            if isinstance(data, list):
                data_text = "\n".join(str(part) for part in data if str(part).strip())
            else:
                data_text = str(data)
            set_cell_text(row.cells[0], str(item.get("no") or index), align="center")
            set_cell_text(row.cells[1], str(item.get("item_name") or ""), align="left")
            set_cell_text(row.cells[2], data_text, align="left")
            set_cell_text(row.cells[3], str(item.get("expected_result") or ""), align="left")
            set_cell_text(row.cells[4], str(item.get("test_result") or ""), align="center")

    current = grid_widths(table._tbl)
    if len(current) >= 5 and target_width and target_width < sum(current):
        set_tbl_width(table._tbl, scaled_widths(current, target_width))
    elif current:
        set_tbl_width(table._tbl, current)


def set_paragraph_bold(p_elm: etree._Element, bold: bool = True) -> None:
    for r_elm in p_elm.findall(qn("w:r")):
        r_pr = ensure_child(r_elm, "w:rPr")
        b = ensure_child(r_pr, "w:b")
        b_cs = ensure_child(r_pr, "w:bCs")
        if bold:
            b.attrib.pop(qn("w:val"), None)
            b_cs.attrib.pop(qn("w:val"), None)
        else:
            b.set(qn("w:val"), "0")
            b_cs.set(qn("w:val"), "0")


def apply_global_yahei_font(doc: Document) -> None:
    for style in doc.styles:
        if hasattr(style, "font") and style.font is not None:
            try:
                style.font.name = FONT_NAME
            except Exception:
                pass
        if hasattr(style, "_element") and style._element is not None:
            r_pr = style._element.find(qn("w:rPr"))
            if r_pr is not None:
                apply_run_font(r_pr)

    body = doc.element.body
    for r_elm in body.iter(qn("w:r")):
        r_pr = ensure_child(r_elm, "w:rPr")
        apply_run_font(r_pr)


def add_paragraph_after(doc: Document, after: etree._Element, text: str, style: str) -> etree._Element:
    paragraph = doc.add_paragraph(text, style=style)
    elm = paragraph._element
    for run_elm in elm.findall(qn("w:r")):
        apply_run_font(ensure_child(run_elm, "w:rPr"))
    after.addnext(elm)
    return elm


def add_picture_after(doc: Document, after: etree._Element, image_path: Path) -> etree._Element:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["FDS內文"]
    run = paragraph.add_run()
    apply_run_font(ensure_child(run._r, "w:rPr"))
    run.add_picture(str(image_path), width=IMAGE_WIDTH)
    elm = paragraph._element
    after.addnext(elm)
    return elm


def clean_step_description(description: str) -> str:
    cleaned = str(description or "").strip()
    cleaned = re.sub(r"^(?:(?:\(\d+\)|\d+[\.、\:\-\s])\s*)+", "", cleaned).strip()
    return cleaned


def fill_attachments(doc: Document, attach_heading: etree._Element, report: dict[str, Any]) -> None:
    set_paragraph_bold(attach_heading, bold=True)
    cursor = attach_heading
    steps = report.get("test_steps") or []
    if not steps:
        cursor = add_paragraph_after(doc, cursor, "（無測試步驟）", "Heading 3")
        return
    for step in steps:
        description = clean_step_description(step.get("description") or "")
        title = description if description else str(step.get("no") or "")
        cursor = add_paragraph_after(doc, cursor, title, "Heading 3")
        images = step.get("images") or []
        if not images:
            continue
        for src in images:
            image_path = ADO_DIR / str(src)
            if not image_path.exists():
                log(f"  找不到圖片，略過：{image_path}")
                continue
            cursor = add_picture_after(doc, cursor, image_path)


def format_history_date(raw: str) -> str:
    """9999-01-01T… → （最新）；正常 ISO → YYYY/M/D HH:MM"""
    text = (raw or "").strip()
    if text.startswith("9999"):
        return "（最新）"
    match = re.match(r"(\d{4})-(\d{2})-(\d{2})T(\d{2}):(\d{2})", text)
    if match:
        y, m, d, hh, mm = match.groups()
        return f"{y}/{int(m)}/{int(d)} {hh}:{mm}"
    return text


def fill_change_history(doc: Document, heading_elm: etree._Element, report: dict[str, Any]) -> None:
    set_paragraph_bold(heading_elm, bold=True)
    history = report.get("change_history") or []
    if not history:
        add_paragraph_after(doc, heading_elm, "（無變動紀錄）", "Heading 3")
        return
    from docx.table import Table as DocxTable

    tbl_elm = OxmlElement("w:tbl")
    tbl_pr = OxmlElement("w:tblPr")
    tbl_elm.append(tbl_pr)
    style = OxmlElement("w:tblStyle")
    style.set(qn("w:val"), "a3")
    tbl_pr.append(style)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "0")
    tbl_w.set(qn("w:type"), "auto")
    tbl_pr.append(tbl_w)
    look = OxmlElement("w:tblLook")
    look.set(qn("w:val"), "04A0")
    tbl_pr.append(look)

    heading_elm.addnext(tbl_elm)
    table = DocxTable(tbl_elm, doc)

    headers = ["時間", "操作者", "欄位", "原值", "新值"]
    total_rows = 1 + len(history)
    for _ in range(total_rows):
        row_elm = OxmlElement("w:tr")
        for _ in range(len(headers)):
            tc = OxmlElement("w:tc")
            p = OxmlElement("w:p")
            tc.append(p)
            row_elm.append(tc)
        tbl_elm.append(row_elm)

    table = DocxTable(tbl_elm, doc)
    for col_idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[col_idx], header, align="center", font_size=FONT_SIZE)
    for row_idx, entry in enumerate(history, start=1):
        row = table.rows[row_idx]
        set_cell_text(row.cells[0], format_history_date(entry.get("date", "")), align="left", font_size=FONT_SIZE)
        set_cell_text(row.cells[1], entry.get("by", ""), align="left", font_size=FONT_SIZE)
        set_cell_text(row.cells[2], entry.get("field", ""), align="left", font_size=FONT_SIZE)
        set_cell_text(row.cells[3], entry.get("old", ""), align="left", font_size=FONT_SIZE)
        set_cell_text(row.cells[4], entry.get("new", ""), align="left", font_size=FONT_SIZE)

    col_widths = [2000, 1800, 1400, 2200, 2200]
    set_tbl_width(tbl_elm, col_widths)


def ado_heading_text(report: dict[str, Any]) -> str:
    ado_ids = report.get("ado_ids") or [report.get("ado_id")]
    ado_text = ", ".join(f"#{value}" for value in ado_ids if value)
    heading = f"ADO{ado_text}"
    redmine = str(report.get("redmine_id") or "").strip()
    if redmine:
        heading += f" / Redmine#{redmine}"
    return heading


def sort_report_ids(report: dict[str, Any]) -> dict[str, Any]:
    """把報告內的 ADO 編號由小到大排好，主旨也跟著排。"""
    raw_ids = report.get("ado_ids") or [report.get("ado_id")]
    ado_ids = sorted(int(value) for value in raw_ids if value is not None)
    ordered = dict(report)
    ordered["ado_ids"] = ado_ids
    if ado_ids:
        ordered["ado_id"] = ado_ids[0]

    lines = [line for line in str(report.get("ado_subject") or "").split("\n") if line.strip()]
    if len(lines) > 1:
        def line_key(line: str) -> int:
            match = re.match(r"^#(\d+)", line.strip())
            return int(match.group(1)) if match else 10**9

        ordered["ado_subject"] = "\n".join(sorted(lines, key=line_key))
    return ordered


def find_sectpr(body: etree._Element) -> etree._Element | None:
    for child in body:
        if child.tag == qn("w:sectPr"):
            return child
    return None


def extract_block_prototype(body: etree._Element) -> list[etree._Element]:
    """取範本的 ADO#XXXX 區塊當樣板（含 meta／測試項目／測試附件）。"""
    placeholder = find_paragraph(body, lambda text: "ADO#XXXX" in text)
    if placeholder is None:
        raise RuntimeError("範本找不到 ADO#XXXX 區塊")
    proto: list[etree._Element] = []
    cur: etree._Element | None = placeholder
    while cur is not None and cur.tag != qn("w:sectPr"):
        proto.append(deepcopy(cur))
        cur = cur.getnext()
    return proto


def clear_ado_blocks(body: etree._Element) -> None:
    """刪掉範本原有的範例與佔位 ADO 區塊。"""
    sample = find_paragraph(body, lambda text: text.startswith("ADO#") and "XXXX" not in text)
    placeholder = find_paragraph(body, lambda text: "ADO#XXXX" in text)
    start = sample if sample is not None else placeholder
    if start is None:
        raise RuntimeError("範本找不到 ADO 區塊")
    cur: etree._Element | None = start
    while cur is not None and cur.tag != qn("w:sectPr"):
        nxt = cur.getnext()
        parent = cur.getparent()
        if parent is not None:
            parent.remove(cur)
        cur = nxt


def append_block(body: etree._Element, proto: list[etree._Element]) -> list[etree._Element]:
    sectpr = find_sectpr(body)
    block = [deepcopy(elm) for elm in proto]
    for elm in block:
        if sectpr is not None:
            sectpr.addprevious(elm)
        else:
            body.append(elm)
    return block


def block_paragraph(block: list[etree._Element], text: str) -> etree._Element | None:
    for elm in block:
        if elm.tag == qn("w:p") and elm_text(elm).strip() == text:
            return elm
    return None


def fill_block(doc: Document, block: list[etree._Element], report: dict[str, Any]) -> None:
    from docx.table import Table

    set_paragraph_text(block[0], ado_heading_text(report))

    meta_tbl = next((elm for elm in block if elm.tag == qn("w:tbl")), None)
    items_heading = block_paragraph(block, "測試項目")
    items_tbl = None
    if items_heading is not None:
        start = block.index(items_heading)
        items_tbl = next((elm for elm in block[start:] if elm.tag == qn("w:tbl")), None)
    attach_heading = block_paragraph(block, "測試附件")
    if meta_tbl is None or items_tbl is None or attach_heading is None:
        raise RuntimeError("範本 ADO 區塊缺少表格或「測試附件」標題")

    fill_meta_table(Table(meta_tbl, doc), report)
    fill_items_table(
        Table(items_tbl, doc),
        list(report.get("test_items") or []),
        target_width=sum(grid_widths(meta_tbl)) or None,
    )

    for elm in block[block.index(attach_heading) + 1 :]:
        parent = elm.getparent()
        if parent is not None:
            parent.remove(elm)
    fill_attachments(doc, attach_heading, report)

    last_elm = attach_heading
    while last_elm.getnext() is not None and last_elm.getnext().tag != qn("w:sectPr"):
        candidate = last_elm.getnext()
        if candidate.tag == qn("w:p") and paragraph_style(candidate) in {"2", "4"}:
            break
        last_elm = candidate
    history_heading = add_paragraph_after(doc, last_elm, "變動紀錄", "Heading 2")
    p_pr = ensure_child(history_heading, "w:pPr")
    p_style = ensure_child(p_pr, "w:pStyle")
    p_style.set(qn("w:val"), "4")
    fill_change_history(doc, history_heading, report)


def paragraph_style(p_elm: etree._Element) -> str:
    p_pr = p_elm.find(qn("w:pPr"))
    if p_pr is None:
        return ""
    p_style = p_pr.find(qn("w:pStyle"))
    return p_style.get(qn("w:val")) if p_style is not None else ""


def max_bookmark_id(root: etree._Element) -> int:
    ids = []
    for node in root.iter(qn("w:bookmarkStart")):
        raw = node.get(qn("w:id")) or ""
        if raw and str(raw).isdigit():
            ids.append(int(raw))
    return max(ids, default=0)


def replace_toc_bookmark(p_elm: etree._Element, bookmark: str) -> None:
    for instr in p_elm.iter(qn("w:instrText")):
        text = instr.text or ""
        if "_Toc" in text:
            instr.text = re.sub(r"_Toc\d+", bookmark, text)


def set_toc_texts(p_elm: etree._Element, number: str, title: str) -> None:
    nodes = list(p_elm.iter(qn("w:t")))
    if not nodes:
        return
    nodes[0].text = f"{number} "
    if len(nodes) > 1:
        nodes[1].text = title


def set_heading_bookmark(p_elm: etree._Element, name: str, bookmark_id: int) -> None:
    for start in list(p_elm.iter(qn("w:bookmarkStart"))):
        old_name = start.get(qn("w:name")) or ""
        if old_name.startswith("_Toc") or old_name == name:
            old_id = start.get(qn("w:id"))
            parent = start.getparent()
            if parent is not None:
                parent.remove(start)
            for end in list(p_elm.iter(qn("w:bookmarkEnd"))):
                if end.get(qn("w:id")) == old_id:
                    end_parent = end.getparent()
                    if end_parent is not None:
                        end_parent.remove(end)
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p_elm.insert(0, start)
    p_elm.append(end)


def collect_ado_outline(body: etree._Element) -> list[dict[str, Any]]:
    outline: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for child in body:
        if child.tag != qn("w:p"):
            continue
        style = paragraph_style(child)
        text = elm_text(child).strip()
        if style == "2" and text.startswith("ADO"):
            current = {"heading": child, "title": text, "items": None, "attach": None, "history": None}
            outline.append(current)
        elif current is not None and style == "4" and text == "測試項目":
            current["items"] = child
        elif current is not None and style == "4" and text == "測試附件":
            current["attach"] = child
        elif current is not None and style == "4" and text == "變動紀錄":
            current["history"] = child
    return outline


def refresh_toc(doc: Document) -> None:
    """依實際 ADO 標題重寫目錄（標題沒有 Redmine 就不出現）。"""
    body = doc.element.body
    sdt = body.find(qn("w:sdt"))
    if sdt is None:
        return
    content = sdt.find(qn("w:sdtContent"))
    if content is None:
        return

    outline = collect_ado_outline(body)
    toc_paras = list(content)
    proto_by_kind: dict[str, etree._Element] = {}
    ado_toc_paras: list[etree._Element] = []
    for para in toc_paras:
        text = elm_text(para)
        if "ADO#" in text:
            proto_by_kind.setdefault("heading", para)
            ado_toc_paras.append(para)
        elif ado_toc_paras and "測試項目" in text:
            proto_by_kind.setdefault("items", para)
            ado_toc_paras.append(para)
        elif ado_toc_paras and "測試附件" in text:
            proto_by_kind.setdefault("attach", para)
            ado_toc_paras.append(para)
        elif ado_toc_paras and "變動紀錄" in text:
            proto_by_kind.setdefault("history", para)
            ado_toc_paras.append(para)

    if "heading" not in proto_by_kind:
        return

    insert_before = ado_toc_paras[-1].getnext() if ado_toc_paras else None
    for para in ado_toc_paras:
        parent = para.getparent()
        if parent is not None:
            parent.remove(para)

    next_id = max_bookmark_id(doc.element) + 1
    for index, block in enumerate(outline, start=2):
        entries = [
            ("heading", f"{index}", block["title"], block["heading"]),
            ("items", f"{index}.1", "測試項目", block.get("items")),
            ("attach", f"{index}.2", "測試附件", block.get("attach")),
            ("history", f"{index}.3", "變動紀錄", block.get("history")),
        ]
        for kind, number, title, heading_elm in entries:
            proto = proto_by_kind.get(kind)
            if proto is None or heading_elm is None:
                continue
            bookmark = f"_Toc{next_id}"
            set_heading_bookmark(heading_elm, bookmark, next_id)
            next_id += 1
            toc_para = deepcopy(proto)
            set_toc_texts(toc_para, number, title)
            replace_toc_bookmark(toc_para, bookmark)
            if insert_before is not None:
                insert_before.addprevious(toc_para)
            else:
                content.append(toc_para)

    for para in list(content):
        nodes = list(para.iter(qn("w:t")))
        if not nodes:
            continue
        number = (nodes[0].text or "").strip()
        if re.fullmatch(r"\d+(?:\.\d+)*", number) and not (nodes[0].text or "").endswith(" "):
            nodes[0].text = f"{number} "


def ensure_heading_number_space(doc: Document) -> None:
    """標題編號與文字之間用空格分隔，避免 11ADO#846。"""
    numbering_part = getattr(doc.part, "numbering_part", None)
    if numbering_part is None:
        return
    heading_styles = {"2", "4", "5", "6", "7"}
    for lvl in numbering_part.element.iter(qn("w:lvl")):
        pstyle = lvl.find(qn("w:pStyle"))
        if pstyle is None or pstyle.get(qn("w:val")) not in heading_styles:
            continue
        suff = lvl.find(qn("w:suff"))
        if suff is None:
            suff = OxmlElement("w:suff")
            lvltext = lvl.find(qn("w:lvlText"))
            if lvltext is not None:
                lvltext.addnext(suff)
            else:
                lvl.append(suff)
        suff.set(qn("w:val"), "space")


def fill_document(doc: Document, reports: list[dict[str, Any]]) -> None:
    ensure_heading_number_space(doc)
    body = doc.element.body
    proto = extract_block_prototype(body)
    clear_ado_blocks(body)
    for report in reports:
        block = append_block(body, proto)
        fill_block(doc, block, report)
    apply_global_yahei_font(doc)
    refresh_toc(doc)


def output_path() -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return DOCX_DIR / f"report_{stamp}.docx"


def collect_reports(work_item_ids: list[int]) -> list[dict[str, Any]]:
    reports: list[dict[str, Any]] = []
    seen_groups: set[tuple[int, ...]] = set()
    for work_item_id in work_item_ids:
        report = sort_report_ids(load_report(work_item_id))
        group = tuple(report.get("ado_ids") or [report.get("ado_id")])
        if group in seen_groups:
            continue
        seen_groups.add(group)
        reports.append(report)
    reports.sort(key=lambda report: (report.get("ado_ids") or [0])[0])
    return reports


def run_word(work_item_ids: list[int]) -> list[str]:
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    log_stage(f"5/5 WORD  產出 Word（{len(work_item_ids)} 筆）")
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"找不到範本：{TEMPLATE_PATH}")

    reports = collect_reports(work_item_ids)
    if not reports:
        raise ValueError("沒有可產出的報告 JSON")
    for report in reports:
        label = ", ".join(str(value) for value in report["ado_ids"])
        log(f"  區塊：ADO {label}")

    path = output_path()
    shutil.copy2(TEMPLATE_PATH, path)
    doc = Document(str(path))
    fill_document(doc, reports)
    doc.save(str(path))
    log(f"已寫入：{path}")
    log_stage_done(f"WORD 完成，共 1 份、{len(reports)} 個 ADO 區塊")
    return [str(path)]


def main() -> None:
    write_pipeline_json(run_word(normalize_ids(read_pipeline_json())))


if __name__ == "__main__":
    run_pipeline_main(main)
