#!/usr/bin/env python3
"""從範本複製 ADO 區塊並填入報告。"""

from __future__ import annotations

import re
from copy import deepcopy
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from lxml import etree

from backend.pipeline.word.oxml import (
    FONT_SIZE,
    META_FONT_SIZE,
    META_NO_START,
    add_paragraph_after,
    add_picture_after,
    elm_text,
    ensure_child,
    ensure_row_count,
    find_paragraph,
    grid_widths,
    paragraph_style,
    scaled_widths,
    set_cell_text,
    set_paragraph_bold,
    set_paragraph_text,
    set_tbl_width,
)
from backend.utils.paths import ADO_DIR
from backend.utils.pipeline_io import log
from backend.utils.report_schema import META_FIELDS


def fill_meta_table(table: Any, report: dict[str, Any]) -> None:
    ensure_row_count(table, 1 + len(META_FIELDS))
    header_row = table.rows[0]
    set_cell_text(header_row.cells[0], "No.", align="center", font_size=META_FONT_SIZE)
    set_cell_text(header_row.cells[1], "項目", align="center", font_size=META_FONT_SIZE)
    set_cell_text(header_row.cells[2], "描述", align="center", font_size=META_FONT_SIZE)
    for index, (label, key) in enumerate(META_FIELDS, start=1):
        row = table.rows[index]
        set_cell_text(
            row.cells[0],
            str(META_NO_START + index - 1),
            align="center",
            font_size=META_FONT_SIZE,
        )
        set_cell_text(row.cells[1], label, align="left", font_size=META_FONT_SIZE)
        set_cell_text(
            row.cells[2],
            str(report.get(key) or ""),
            align="left",
            font_size=META_FONT_SIZE,
        )


def fill_items_table(table: Any, items: list[dict[str, Any]], *, target_width: int | None = None) -> None:
    rows_needed = max(2, 1 + len(items))
    ensure_row_count(table, rows_needed)
    if not items:
        set_cell_text(table.rows[1].cells[0], "1", align="center")
        for col in range(1, 5):
            set_cell_text(table.rows[1].cells[col], "", align="left")
    else:
        while len(table.rows) > 1 + len(items):
            table._tbl.remove(table.rows[-1]._tr)
        for index, item in enumerate(items, start=1):
            row = table.rows[index]
            data = item.get("test_data") or []
            if isinstance(data, list):
                data_text = "\n".join(str(part) for part in data if str(part).strip())
            else:
                data_text = str(data)
            set_cell_text(row.cells[0], str(item.get("no") or index), align="center")
            set_cell_text(row.cells[1], str(item.get("item_name") or ""), align="left")
            set_cell_text(row.cells[2], data_text, align="left")
            set_cell_text(row.cells[3], str(item.get("expected_result") or ""), align="left")
            set_cell_text(row.cells[4], str(item.get("test_result") or ""), align="center")

    current = grid_widths(table._tbl)
    if len(current) >= 5 and target_width and target_width < sum(current):
        set_tbl_width(table._tbl, scaled_widths(current, target_width))
    elif current:
        set_tbl_width(table._tbl, current)


def clean_step_description(description: str) -> str:
    cleaned = str(description or "").strip()
    cleaned = re.sub(r"^(?:(?:\(\d+\)|\d+[\.、\:\-\s])\s*)+", "", cleaned).strip()
    return cleaned


def fill_attachments(doc: Document, attach_heading: etree._Element, report: dict[str, Any]) -> None:
    set_paragraph_bold(attach_heading, bold=True)
    cursor = attach_heading
    steps = report.get("test_steps") or []
    if not steps:
        add_paragraph_after(doc, cursor, "（無測試步驟）", "Heading 3")
        return
    for step in steps:
        description = clean_step_description(step.get("description") or "")
        title = description if description else str(step.get("no") or "")
        cursor = add_paragraph_after(doc, cursor, title, "Heading 3")
        images = step.get("images") or []
        if not images:
            continue
        for src in images:
            image_path = ADO_DIR / str(src)
            if not image_path.exists():
                log(f"  找不到圖片，略過：{image_path}")
                continue
            cursor = add_picture_after(doc, cursor, image_path)


def format_commit_date(raw: str) -> str:
    match = re.match(r"(\d{4})-(\d{2})-(\d{2})", (raw or "").strip())
    if match:
        year, month, day = match.groups()
        return f"{year}/{int(month)}/{int(day)}"
    return (raw or "").strip()


def last_block_element(start: etree._Element) -> etree._Element:
    last_elm = start
    while last_elm.getnext() is not None and last_elm.getnext().tag != qn("w:sectPr"):
        candidate = last_elm.getnext()
        if candidate.tag == qn("w:p") and paragraph_style(candidate) in {"2", "4"}:
            break
        last_elm = candidate
    return last_elm


def add_heading4_after(doc: Document, after: etree._Element, text: str) -> etree._Element:
    heading = add_paragraph_after(doc, after, text, "Heading 2")
    p_pr = ensure_child(heading, "w:pPr")
    p_style = ensure_child(p_pr, "w:pStyle")
    p_style.set(qn("w:val"), "4")
    set_paragraph_bold(heading, bold=True)
    return heading


def fill_commits(doc: Document, heading_elm: etree._Element, commits: list[dict[str, Any]]) -> None:
    show_ado = any(entry.get("ado_id") for entry in commits)
    headers = (["ADO"] if show_ado else []) + ["SHA", "訊息", "作者", "日期"]
    tbl_elm = OxmlElement("w:tbl")
    tbl_pr = OxmlElement("w:tblPr")
    tbl_elm.append(tbl_pr)
    style = OxmlElement("w:tblStyle")
    style.set(qn("w:val"), "a3")
    tbl_pr.append(style)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), "0")
    tbl_w.set(qn("w:type"), "auto")
    tbl_pr.append(tbl_w)
    look = OxmlElement("w:tblLook")
    look.set(qn("w:val"), "04A0")
    tbl_pr.append(look)
    heading_elm.addnext(tbl_elm)

    for _ in range(1 + len(commits)):
        row_elm = OxmlElement("w:tr")
        for _ in headers:
            tc = OxmlElement("w:tc")
            tc.append(OxmlElement("w:p"))
            row_elm.append(tc)
        tbl_elm.append(row_elm)

    table = Table(tbl_elm, doc)
    for col_idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[col_idx], header, align="center", font_size=FONT_SIZE)
    for row_idx, entry in enumerate(commits, start=1):
        row = table.rows[row_idx]
        values = []
        if show_ado:
            ado_id = entry.get("ado_id")
            values.append(f"#{ado_id}" if ado_id else "")
        values.extend(
            [
                str(entry.get("sha_short") or (entry.get("sha") or "")[:8]),
                str(entry.get("message") or ""),
                str(entry.get("author") or ""),
                format_commit_date(str(entry.get("date") or "")),
            ]
        )
        for col_idx, value in enumerate(values):
            set_cell_text(row.cells[col_idx], value, align="left", font_size=FONT_SIZE)

    col_widths = ([1200] if show_ado else []) + [1600, 4200, 1800, 1800]
    set_tbl_width(tbl_elm, col_widths)


def ado_heading_text(report: dict[str, Any]) -> str:
    ado_ids = report.get("ado_ids") or [report.get("ado_id")]
    ado_text = ", ".join(f"#{value}" for value in ado_ids if value)
    heading = f"ADO{ado_text}"
    redmine = str(report.get("redmine_id") or "").strip()
    if redmine:
        heading += f" / Redmine#{redmine}"
    return heading


def sort_report_ids(report: dict[str, Any]) -> dict[str, Any]:
    """把報告內的 ADO 編號由小到大排好，主旨也跟著排。"""
    raw_ids = report.get("ado_ids") or [report.get("ado_id")]
    ado_ids = sorted(int(value) for value in raw_ids if value is not None)
    ordered = dict(report)
    ordered["ado_ids"] = ado_ids
    if ado_ids:
        ordered["ado_id"] = ado_ids[0]

    lines = [line for line in str(report.get("ado_subject") or "").split("\n") if line.strip()]
    if len(lines) > 1:
        def line_key(line: str) -> int:
            match = re.match(r"^#(\d+)", line.strip())
            return int(match.group(1)) if match else 10**9

        ordered["ado_subject"] = "\n".join(sorted(lines, key=line_key))
    return ordered


def find_sectpr(body: etree._Element) -> etree._Element | None:
    for child in body:
        if child.tag == qn("w:sectPr"):
            return child
    return None


def extract_block_prototype(body: etree._Element) -> list[etree._Element]:
    """取範本的 ADO#XXXX 區塊當樣板（含 meta／測試項目／測試附件）。"""
    placeholder = find_paragraph(body, lambda text: "ADO#XXXX" in text)
    if placeholder is None:
        raise RuntimeError("範本找不到 ADO#XXXX 區塊")
    proto: list[etree._Element] = []
    cur: etree._Element | None = placeholder
    while cur is not None and cur.tag != qn("w:sectPr"):
        proto.append(deepcopy(cur))
        cur = cur.getnext()
    return proto


def clear_ado_blocks(body: etree._Element) -> None:
    """刪掉範本原有的範例與佔位 ADO 區塊。"""
    sample = find_paragraph(body, lambda text: text.startswith("ADO#") and "XXXX" not in text)
    placeholder = find_paragraph(body, lambda text: "ADO#XXXX" in text)
    start = sample if sample is not None else placeholder
    if start is None:
        raise RuntimeError("範本找不到 ADO 區塊")
    cur: etree._Element | None = start
    while cur is not None and cur.tag != qn("w:sectPr"):
        nxt = cur.getnext()
        parent = cur.getparent()
        if parent is not None:
            parent.remove(cur)
        cur = nxt


def append_block(body: etree._Element, proto: list[etree._Element]) -> list[etree._Element]:
    sectpr = find_sectpr(body)
    block = [deepcopy(elm) for elm in proto]
    for elm in block:
        if sectpr is not None:
            sectpr.addprevious(elm)
        else:
            body.append(elm)
    return block


def block_paragraph(block: list[etree._Element], text: str) -> etree._Element | None:
    for elm in block:
        if elm.tag == qn("w:p") and elm_text(elm).strip() == text:
            return elm
    return None


def fill_block(doc: Document, block: list[etree._Element], report: dict[str, Any]) -> None:
    set_paragraph_text(block[0], ado_heading_text(report))

    meta_tbl = next((elm for elm in block if elm.tag == qn("w:tbl")), None)
    items_heading = block_paragraph(block, "測試項目")
    items_tbl = None
    if items_heading is not None:
        start = block.index(items_heading)
        items_tbl = next((elm for elm in block[start:] if elm.tag == qn("w:tbl")), None)
    attach_heading = block_paragraph(block, "測試附件")
    if meta_tbl is None or items_tbl is None or attach_heading is None:
        raise RuntimeError("範本 ADO 區塊缺少表格或「測試附件」標題")

    fill_meta_table(Table(meta_tbl, doc), report)
    fill_items_table(
        Table(items_tbl, doc),
        list(report.get("test_items") or []),
        target_width=sum(grid_widths(meta_tbl)) or None,
    )

    for elm in block[block.index(attach_heading) + 1 :]:
        parent = elm.getparent()
        if parent is not None:
            parent.remove(elm)
    fill_attachments(doc, attach_heading, report)
    commits = [entry for entry in (report.get("commits") or []) if isinstance(entry, dict)]
    if commits:
        heading = add_heading4_after(doc, last_block_element(attach_heading), "Commit")
        fill_commits(doc, heading, commits)
