#!/usr/bin/env python3
"""Word 文件的 OOXML 小工具。"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

IMAGE_WIDTH = Inches(6.2)
FONT_NAME = "Microsoft YaHei"
FONT_SIZE = "16"
META_FONT_SIZE = "20"
META_NO_START = 1


def elm_text(elm: etree._Element) -> str:
    return "".join(node.text or "" for node in elm.iter(qn("w:t")))


def set_paragraph_text(p_elm: etree._Element, text: str) -> None:
    nodes = list(p_elm.iter(qn("w:t")))
    if not nodes:
        r = p_elm.makeelement(qn("w:r"), {})
        t = r.makeelement(qn("w:t"), {})
        t.text = text
        p_elm.append(r)
        return
    nodes[0].text = text
    if text.startswith(" ") or text.endswith(" "):
        nodes[0].set(qn("xml:space"), "preserve")
    for node in nodes[1:]:
        node.text = ""


def ensure_child(parent: etree._Element, tag: str) -> etree._Element:
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def apply_run_font(r_pr: etree._Element, font_size: str | None = None) -> None:
    r_fonts = ensure_child(r_pr, "w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT_NAME)
    r_fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_fonts.set(qn("w:eastAsia"), FONT_NAME)
    r_fonts.set(qn("w:cs"), FONT_NAME)
    r_fonts.set(qn("w:hint"), "eastAsia")
    if font_size is not None:
        ensure_child(r_pr, "w:sz").set(qn("w:val"), font_size)
        ensure_child(r_pr, "w:szCs").set(qn("w:val"), font_size)
    lang = ensure_child(r_pr, "w:lang")
    lang.set(qn("w:eastAsia"), "zh-TW")


def clean_cell_format(cell: Any, align: str = "left", font_size: str = FONT_SIZE) -> None:
    tc_pr = ensure_child(cell._tc, "w:tcPr")
    no_wrap = tc_pr.find(qn("w:noWrap"))
    if no_wrap is not None:
        tc_pr.remove(no_wrap)
    ensure_child(tc_pr, "w:vAlign").set(qn("w:val"), "center")

    for paragraph in cell.paragraphs:
        p_pr = ensure_child(paragraph._p, "w:pPr")
        p_style = p_pr.find(qn("w:pStyle"))
        if p_style is not None:
            p_pr.remove(p_style)
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is not None:
            p_pr.remove(num_pr)
        indent = p_pr.find(qn("w:ind"))
        if indent is not None:
            p_pr.remove(indent)
        ensure_child(p_pr, "w:jc").set(qn("w:val"), align)
        apply_run_font(ensure_child(p_pr, "w:rPr"), font_size=font_size)


def set_cell_text(
    cell: Any,
    text: str,
    align: str = "left",
    font_size: str = FONT_SIZE,
) -> None:
    lines = [line for line in str(text or "").split("\n")]
    if not lines:
        lines = [""]
    while len(cell.paragraphs) > 1:
        cell.paragraphs[-1]._element.getparent().remove(cell.paragraphs[-1]._element)

    def write_paragraph(paragraph: Any, line: str) -> None:
        proto = deepcopy(paragraph.runs[0]._r) if paragraph.runs else None
        for run in list(paragraph.runs):
            run._element.getparent().remove(run._element)
        if proto is None:
            run_elm = OxmlElement("w:r")
            paragraph._p.append(run_elm)
        else:
            run_elm = deepcopy(proto)
            for child in list(run_elm):
                if child.tag in {qn("w:t"), qn("w:br"), qn("w:tab")}:
                    run_elm.remove(child)
            paragraph._p.append(run_elm)
        r_pr = ensure_child(run_elm, "w:rPr")
        apply_run_font(r_pr, font_size=font_size)
        t_elm = OxmlElement("w:t")
        t_elm.text = line
        if line[:1].isspace() or line[-1:].isspace():
            t_elm.set(qn("xml:space"), "preserve")
        run_elm.append(t_elm)

    write_paragraph(cell.paragraphs[0], lines[0])
    for line in lines[1:]:
        new_p = deepcopy(cell.paragraphs[0]._p)
        for run in list(new_p.findall(qn("w:r"))):
            new_p.remove(run)
        cell._tc.append(new_p)
        write_paragraph(cell.paragraphs[-1], line)
    clean_cell_format(cell, align=align, font_size=font_size)


def find_paragraph(body: etree._Element, predicate) -> etree._Element | None:
    for child in list(body):
        if child.tag == qn("w:p") and predicate(elm_text(child)):
            return child
    return None


def ensure_row_count(table: Any, total_rows: int) -> None:
    tbl = table._tbl
    while len(table.rows) < total_rows:
        tbl.append(deepcopy(tbl.tr_lst[-1]))
    while len(table.rows) > total_rows:
        tbl.remove(tbl.tr_lst[-1])


def set_tbl_width(tbl: etree._Element, widths: list[int]) -> None:
    """固定表格與欄寬（dxa），並取消負縮排。"""
    total = sum(widths)
    tbl_pr = ensure_child(tbl, "w:tblPr")
    tbl_w = ensure_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is not None:
        tbl_ind.set(qn("w:w"), "0")
        tbl_ind.set(qn("w:type"), "dxa")
    layout = ensure_child(tbl_pr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(1, grid)
    cols = list(grid.findall(qn("w:gridCol")))
    while len(cols) < len(widths):
        col = OxmlElement("w:gridCol")
        grid.append(col)
        cols.append(col)
    for col in cols[len(widths) :]:
        grid.remove(col)
    for col, width in zip(cols, widths):
        col.set(qn("w:w"), str(width))

    for row in tbl.findall(qn("w:tr")):
        cells = row.findall(qn("w:tc"))
        for cell, width in zip(cells, widths):
            tc_pr = ensure_child(cell, "w:tcPr")
            tc_w = ensure_child(tc_pr, "w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def grid_widths(tbl: etree._Element) -> list[int]:
    grid = tbl.find(qn("w:tblGrid"))
    if grid is None:
        return []
    widths: list[int] = []
    for col in grid.findall(qn("w:gridCol")):
        raw = col.get(qn("w:w")) or "0"
        widths.append(int(raw) if str(raw).isdigit() else 0)
    return widths


def scaled_widths(source: list[int], target_total: int) -> list[int]:
    current = sum(source) or 1
    scaled = [max(1, round(width * target_total / current)) for width in source]
    scaled[-1] += target_total - sum(scaled)
    return scaled


def set_paragraph_bold(p_elm: etree._Element, bold: bool = True) -> None:
    for r_elm in p_elm.findall(qn("w:r")):
        r_pr = ensure_child(r_elm, "w:rPr")
        b = ensure_child(r_pr, "w:b")
        b_cs = ensure_child(r_pr, "w:bCs")
        if bold:
            b.attrib.pop(qn("w:val"), None)
            b_cs.attrib.pop(qn("w:val"), None)
        else:
            b.set(qn("w:val"), "0")
            b_cs.set(qn("w:val"), "0")


def apply_global_yahei_font(doc: Document) -> None:
    for style in doc.styles:
        if hasattr(style, "font") and style.font is not None:
            try:
                style.font.name = FONT_NAME
            except Exception:
                pass
        if hasattr(style, "_element") and style._element is not None:
            r_pr = style._element.find(qn("w:rPr"))
            if r_pr is not None:
                apply_run_font(r_pr)

    body = doc.element.body
    for r_elm in body.iter(qn("w:r")):
        r_pr = ensure_child(r_elm, "w:rPr")
        apply_run_font(r_pr)


def add_paragraph_after(doc: Document, after: etree._Element, text: str, style: str) -> etree._Element:
    paragraph = doc.add_paragraph(text, style=style)
    elm = paragraph._element
    for run_elm in elm.findall(qn("w:r")):
        apply_run_font(ensure_child(run_elm, "w:rPr"))
    after.addnext(elm)
    return elm


def add_picture_after(doc: Document, after: etree._Element, image_path: Path) -> etree._Element:
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["FDS內文"]
    run = paragraph.add_run()
    apply_run_font(ensure_child(run._r, "w:rPr"))
    with image_path.open("rb") as handle:
        run.add_picture(handle, width=IMAGE_WIDTH)
    elm = paragraph._element
    after.addnext(elm)
    return elm


def paragraph_style(p_elm: etree._Element) -> str:
    p_pr = p_elm.find(qn("w:pPr"))
    if p_pr is None:
        return ""
    p_style = p_pr.find(qn("w:pStyle"))
    return p_style.get(qn("w:val")) if p_style is not None else ""
