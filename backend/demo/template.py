#!/usr/bin/env python3
"""展示用 Word 範本：結構對齊填表程式（ADO#XXXX、測試項目、測試附件、FDS內文）。"""

from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from backend.utils.paths import DEMO_TEMPLATE_PATH

META_HEADERS = ("No.", "項目", "描述")
META_LABELS = ("ADO主旨", "測試環境", "Site", "版本號", "測試日期", "測試人員", "審核人員")
ITEM_HEADERS = ("No.", "測試項目", "測試資料", "預期結果", "測試結果")


def _set_p_style(paragraph, style_id: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_style = p_pr.get_or_add_pStyle()
    p_style.set(qn("w:val"), style_id)


def _ensure_style_id(doc: Document, style_id: str, display_name: str, based_on: str) -> None:
    styles_elm = doc.styles.element
    for style in styles_elm.findall(qn("w:style")):
        if style.get(qn("w:styleId")) == style_id:
            return
    clone = deepcopy(doc.styles[based_on]._element)
    clone.set(qn("w:styleId"), style_id)
    clone.attrib.pop(qn("w:default"), None)
    name = clone.find(qn("w:name"))
    if name is not None:
        name.set(qn("w:val"), display_name)
    styles_elm.append(clone)


def _ensure_body_style(doc: Document) -> None:
    try:
        doc.styles["FDS內文"]
        return
    except KeyError:
        pass
    style = doc.styles.add_style("FDS內文", WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(12)


def _fill_row(table, row_idx: int, values: tuple[str, ...]) -> None:
    for col, value in enumerate(values):
        table.rows[row_idx].cells[col].text = value


def _add_ado_block(doc: Document, heading: str) -> None:
    title = doc.add_heading(heading, level=1)
    _set_p_style(title, "2")

    meta = doc.add_table(rows=1 + len(META_LABELS), cols=3)
    meta.style = "Table Grid"
    _fill_row(meta, 0, META_HEADERS)
    for index, label in enumerate(META_LABELS, start=1):
        _fill_row(meta, index, (str(index), label, ""))

    items_heading = doc.add_heading("測試項目", level=3)
    _set_p_style(items_heading, "4")
    items = doc.add_table(rows=2, cols=5)
    items.style = "Table Grid"
    _fill_row(items, 0, ITEM_HEADERS)
    _fill_row(items, 1, ("1", "", "", "", ""))

    attach = doc.add_heading("測試附件", level=3)
    _set_p_style(attach, "4")
    doc.add_paragraph("（測試步驟與截圖會寫在此處）")


def build_demo_template(path: Path | None = None) -> Path:
    dest = path or DEMO_TEMPLATE_PATH
    dest.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _ensure_style_id(doc, "2", "展示標題1", "Heading 1")
    _ensure_style_id(doc, "4", "展示標題3", "Heading 3")
    _ensure_body_style(doc)

    cover = doc.add_paragraph("測試報告")
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.runs[0].font.size = Pt(28)
    cover.runs[0].font.color.rgb = RGBColor(0x16, 0x16, 0x16)
    cover.runs[0].bold = True

    subtitle = doc.add_paragraph("展示用範本（虛構專案 LumenWMS，不含客戶資料）")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(0x52, 0x52, 0x52)

    doc.add_heading("目錄", level=1)
    toc = doc.add_paragraph("（產出 Word 時會依 ADO 區塊更新目錄；此處僅為佔位。）")
    toc.runs[0].italic = True

    doc.add_heading("測試紀錄", level=1)
    _add_ado_block(doc, "ADO#8401 / Redmine#64001")
    _add_ado_block(doc, "ADO#XXXX")

    doc.save(dest)
    return dest
